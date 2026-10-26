# -*- coding: utf-8 -*-
# SPDX-License-Identifier: MIT
"""Tests fuer das claude_code-Backend der Foerderbericht-Pipeline (_call_llm)."""

import subprocess
import sys
import types
from pathlib import Path

import pytest

SYSTEM_ROOT = Path(__file__).parent.parent
if str(SYSTEM_ROOT) not in sys.path:
    sys.path.insert(0, str(SYSTEM_ROOT))

from hub._services.document import foerderbericht_pipeline as fp
from hub._services.document.foerderbericht_pipeline import FoerderberichtPipeline, PipelineError


def _pipeline(tmp_path):
    return FoerderberichtPipeline(base_path=tmp_path)


def test_call_llm_passes_prompt_via_stdin_without_tools(tmp_path, monkeypatch):
    seen = {}

    def fake_run(cmd, **kwargs):
        seen["cmd"], seen["kwargs"] = cmd, kwargs
        return types.SimpleNamespace(returncode=0, stdout=' {"ok": 1} \n', stderr="")

    monkeypatch.setattr(fp.subprocess, "run", fake_run)
    big_prompt = "x" * 100_000  # ueber dem Windows-argv-Limit

    out = _pipeline(tmp_path)._call_llm(big_prompt, "claude_code", "claude-sonnet-4-6")

    assert out == '{"ok": 1}'
    assert seen["kwargs"]["input"] == big_prompt
    assert big_prompt not in seen["cmd"]
    cmd = seen["cmd"]
    assert cmd[cmd.index("--tools") + 1] == ""  # keine Tools
    assert "--allowedTools" not in cmd
    assert seen["kwargs"]["env"]["CLAUDE_CODE_DISABLE_CLAUDE_MDS"] == "1"
    assert Path(seen["kwargs"]["cwd"]) != Path.cwd()


def test_call_llm_works_when_system_core_package_is_loaded(tmp_path, monkeypatch):
    """Regression: 'from core.runner import ClaudeRunner' scheiterte, weil
    system/core das Paket 'core' belegt."""
    import core  # noqa: F401  (system/core)

    monkeypatch.setattr(
        fp.subprocess, "run",
        lambda cmd, **kw: types.SimpleNamespace(returncode=0, stdout="{}", stderr=""),
    )
    assert _pipeline(tmp_path)._call_llm("p", "claude_code", "m") == "{}"


def test_call_llm_reports_cli_errors(tmp_path, monkeypatch):
    monkeypatch.setattr(
        fp.subprocess, "run",
        lambda cmd, **kw: types.SimpleNamespace(returncode=1, stdout="", stderr="auth fehlt"),
    )
    with pytest.raises(PipelineError, match="auth fehlt"):
        _pipeline(tmp_path)._call_llm("p", "claude_code", "m")

    def missing(cmd, **kw):
        raise FileNotFoundError

    monkeypatch.setattr(fp.subprocess, "run", missing)
    with pytest.raises(PipelineError, match="PATH"):
        _pipeline(tmp_path)._call_llm("p", "llmauto", "m")

    def slow(cmd, **kw):
        raise subprocess.TimeoutExpired(cmd, 600)

    monkeypatch.setattr(fp.subprocess, "run", slow)
    with pytest.raises(PipelineError, match="Timeout"):
        _pipeline(tmp_path)._call_llm("p", "claude_code", "m")


def test_replace_placeholders_treats_none_as_empty_text():
    """Regression: 'null' in der LLM-Antwort (zeichen_amt, sachbearbeiter, ...) liess
    fill_template mit TypeError scheitern; die Pipeline fiel still auf den
    Markdown-Generator zurueck."""
    from docx import Document
    from hub._services.document.word_template_service import WordTemplateService

    doc = Document()
    doc.add_paragraph("Aktenzeichen: {{ZEICHEN}} Ende")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Sachbearbeiter {{SB}}"

    WordTemplateService().replace_placeholders(doc, {"{{ZEICHEN}}": None, "{{SB}}": None})

    assert doc.paragraphs[0].text == "Aktenzeichen:  Ende"
    assert table.cell(0, 0).text == "Sachbearbeiter "
