from __future__ import annotations

import hashlib
import json
import secrets
from pathlib import Path

import pytest

from hub._services.document.report_provider_adapter import (
    BachOperatorConfirmationPort,
    BachDocumentAnonymizerPort,
    BachDocumentReaderPort,
    BachModelRoutePort,
    BachWordTemplatePort,
    ReportHostPortError,
    ReportProviderAdapter,
    build_report_host_ports,
    load_report_provider,
)
from hub._services.document.report_provider_contract import (
    PROVIDER_REASON_CODES,
    SCHEMA,
    PreparedRun,
    ProviderCapabilities,
    ReportContractError,
    ReportRequest,
    ReportResult,
    WatcherConfig,
)


_SUBJECT_REF = "sub_" + "1" * 48
_SOURCE_REF = "src_" + "2" * 48
_SOURCE_REVISION = hashlib.sha256(b"synthetic source bytes").hexdigest()
_TEMPLATE_REVISION = "template_revision_1"
_DOCX_MEDIA = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_CAPABILITIES = {
    "schema": SCHEMA,
    "available": True,
    "modes": ["local"],
    "watcher_available": True,
    "mode_statuses": {
        "local": {"available": True},
        "cloud": {"available": False, "reason_code": "cloud_consent_unavailable"},
        "hybrid": {"available": False, "reason_code": "hybrid_policy_unavailable"},
    },
}


def _request_wire():
    from datetime import datetime, timedelta, timezone

    stable = {
        "schema": SCHEMA,
        "request_id": "request-opaque-1",
        "subject_ref": _SUBJECT_REF,
        "source_refs": [_SOURCE_REF],
        "period_start": "2026-01-01",
        "period_end": "2026-01-31",
        "mode": "local",
        "template_id": "template-opaque-1",
    }
    digest = hashlib.sha256(
        json.dumps(stable, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    return {
        **stable,
        "mode_consent": {
            "mode": "local",
            "consent_ref": "consent-opaque-1",
            "request_id": stable["request_id"],
            "request_digest": digest,
            "approved": True,
            "expires_at": (datetime.now(timezone.utc) + timedelta(minutes=10)).isoformat(),
        },
    }


def _watcher_wire(*, enabled=True, subject_refs=None):
    from datetime import datetime, timedelta, timezone

    return {
        "schema": SCHEMA,
        "enabled": enabled,
        "auto_finish": False,
        "subject_refs": [_SUBJECT_REF] if subject_refs is None and enabled else (subject_refs or []),
        "mode": "local",
        "template_id": "template-opaque-1",
        "period_start": "2026-01-01",
        "period_end": "2026-01-31",
        "consent_ref": "consent-opaque-1",
        "consent_expires_at": (datetime.now(timezone.utc) + timedelta(minutes=10)).isoformat(),
    }


def test_capabilities_wire_keeps_per_mode_availability_and_neutral_reasons():
    parsed = ProviderCapabilities.from_wire(_CAPABILITIES)

    assert parsed.to_wire() == {
        **_CAPABILITIES,
        "modes": ["local"],
    }


def test_provider_reason_vocabulary_matches_the_normalized_facade_set():
    assert PROVIDER_REASON_CODES == frozenset({
        "prepared", "draft_created", "cleanup_completed", "watcher_configured",
        "watcher_state_read", "watcher_tick_completed", "provider_configured",
        "provider_not_configured", "provider_unavailable", "provider_invalid_request",
        "provider_consent_denied", "provider_consent_expired", "provider_source_unavailable",
        "provider_source_changed", "provider_template_unavailable", "provider_template_changed",
        "provider_run_store_unavailable", "provider_run_conflict", "provider_finish_denied",
        "provider_route_unverified", "provider_anonymization_unsupported",
        "provider_anonymization_receipt_invalid", "provider_hybrid_schema_invalid",
        "provider_report_schema_invalid", "provider_render_unavailable", "provider_effect_unknown",
        "provider_cleanup_denied", "provider_cleanup_unavailable", "provider_cleanup_partial",
        "provider_run_not_found", "provider_watcher_disabled", "provider_watcher_unavailable",
        "provider_watcher_consent_denied",
    })


@pytest.mark.parametrize("update", [
    {"mode_statuses": None},
    {"modes": ["local", "cloud"]},
    {"modes": ["local", "local"]},
    {"available": False},
    {"private_prompt": "must not enter public capabilities"},
    {"mode_statuses": {
        "local": {"available": True},
        "cloud": {"available": False, "reason_code": "private/path"},
        "hybrid": {"available": False, "reason_code": "hybrid_policy_unavailable"},
    }},
    {"mode_statuses": {
        "local": {"available": True},
        "cloud": {"available": False, "reason_code": "provider.source_unavailable"},
        "hybrid": {"available": False, "reason_code": "hybrid_policy_unavailable"},
    }},
])
def test_capabilities_reject_mismatched_or_private_projection(update):
    payload = {**_CAPABILITIES, **update}
    with pytest.raises(ReportContractError):
        ProviderCapabilities.from_wire(payload)


@pytest.mark.asyncio
async def test_unavailable_capabilities_preserve_reason_for_each_mode():
    result = await ReportProviderAdapter.unavailable("provider_missing").capabilities()
    parsed = ProviderCapabilities.from_wire(result)

    assert parsed.available is False
    assert parsed.modes == ()
    assert parsed.reason_code == "provider_missing"
    assert all(
        parsed.mode_statuses[mode]["reason_code"] == "provider_missing"
        for mode in ("local", "cloud", "hybrid")
    )


@pytest.mark.parametrize("subject_refs", [
    [],
    [_SUBJECT_REF, _SUBJECT_REF],
    ["private-person-name"],
    ["sub_" + "a" * 47],
    ["sub_" + format(index, "048x") for index in range(201)],
])
def test_enabled_watcher_requires_bounded_unique_opaque_subject_scope(subject_refs):
    with pytest.raises(ReportContractError):
        WatcherConfig.from_wire(_watcher_wire(subject_refs=subject_refs))


def test_disabled_watcher_can_be_configured_with_empty_subject_scope():
    config = WatcherConfig.from_wire(_watcher_wire(enabled=False, subject_refs=[]))

    assert config.enabled is False
    assert config.subject_refs == ()
    assert config.to_wire()["subject_refs"] == []
    assert config.to_wire()["schema"] == SCHEMA


def test_watcher_rejects_wrong_or_missing_schema():
    for payload in (
        {key: value for key, value in _watcher_wire().items() if key != "schema"},
        {**_watcher_wire(), "schema": "foerderplaner.report.v1"},
    ):
        with pytest.raises(ReportContractError, match="invalid_watcher_schema"):
            WatcherConfig.from_wire(payload)


def test_prepared_run_source_count_matches_request_bound():
    valid = {
        "schema": SCHEMA,
        "run_id": "run-opaque-200",
        "status": "prepared",
        "preview_ref": "preview-opaque-200",
        "expires_at": "2099-01-01T00:00:00+00:00",
        "request_digest": hashlib.sha256(b"request").hexdigest(),
        "source_revision": hashlib.sha256(b"sources").hexdigest(),
        "template_revision": "template_revision_1",
        "source_count": 200,
        "reason_code": None,
    }
    assert PreparedRun.from_wire(valid).source_count == 200
    with pytest.raises(ReportContractError, match="invalid_source_count"):
        PreparedRun.from_wire({**valid, "source_count": 201})


def test_prepared_run_rejects_non_neutral_reason_code():
    blocked = {
        "schema": SCHEMA,
        "run_id": None,
        "status": "blocked",
        "preview_ref": None,
        "expires_at": None,
        "request_digest": hashlib.sha256(b"request").hexdigest(),
        "source_revision": None,
        "template_revision": None,
        "source_count": 0,
        "reason_code": "provider_source_unavailable",
    }
    assert PreparedRun.from_wire(blocked).reason_code == "provider_source_unavailable"
    with pytest.raises(ReportContractError, match="invalid_reason_code"):
        PreparedRun.from_wire({**blocked, "reason_code": "private_directory_name"})
    with pytest.raises(ReportContractError, match="invalid_reason_code"):
        PreparedRun.from_wire({**blocked, "reason_code": "provider.source_unavailable"})


def test_report_result_rejects_non_neutral_reason_code():
    draft = {
        "schema": SCHEMA,
        "run_id": "run-opaque-1",
        "status": "draft",
        "effect": "completed",
        "reason_code": "draft_created",
        "artifact_refs": ["artifact-opaque-1"],
        "review_required": True,
    }
    assert ReportResult.from_wire(draft).reason_code == "draft_created"
    assert ReportResult.from_wire({
        **draft,
        "reason_code": "provider_source_unavailable",
        "status": "blocked",
        "effect": "none",
        "artifact_refs": [],
    }).reason_code == "provider_source_unavailable"
    with pytest.raises(ReportContractError, match="invalid_reason_code"):
        ReportResult.from_wire({**draft, "reason_code": "private_directory_name"})
    with pytest.raises(ReportContractError, match="invalid_reason_code"):
        ReportResult.from_wire({**draft, "reason_code": "provider.source_unavailable"})
_MINIMAL_TEXT_PDF = b"""%PDF-1.4
1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj
2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj
3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 100]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>endobj
4 0 obj<</Length 54>>stream
BT /F1 12 Tf 20 50 Td (Probe Dokument Text) Tj ET
endstream
endobj
5 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj
xref
0 6
trailer<</Size 6/Root 1 0 R>>
startxref
0
%%EOF
"""


class _SourceProvider:
    def __init__(self, payload):
        self.payload = payload

    async def list_sources(self, subject_ref):
        assert subject_ref == _SUBJECT_REF
        return self.payload


@pytest.mark.asyncio
async def test_list_sources_projects_only_complete_opaque_metadata():
    provider = _SourceProvider({
        "schema": SCHEMA,
        "complete": True,
        "sources": [{
            "source_ref": _SOURCE_REF,
            "source_revision": _SOURCE_REVISION,
            "media_type": "text/plain",
        }],
    })

    result = await ReportProviderAdapter(provider).list_sources(_SUBJECT_REF)

    assert result == {
        "schema": SCHEMA,
        "complete": True,
        "sources": [{
            "source_ref": _SOURCE_REF,
            "source_revision": _SOURCE_REVISION,
            "media_type": "text/plain",
        }],
    }


@pytest.mark.asyncio
@pytest.mark.parametrize("payload", [
    {"schema": SCHEMA, "complete": False, "sources": []},
    {"schema": SCHEMA, "complete": False, "sources": [{
        "source_ref": _SOURCE_REF,
        "source_revision": _SOURCE_REVISION,
        "media_type": "text/plain",
    }]},
    {"schema": SCHEMA, "complete": True, "sources": [{
        "source_ref": _SOURCE_REF,
        "source_revision": _SOURCE_REVISION,
        "media_type": "text/plain",
    }] * 201},
    {"schema": SCHEMA, "complete": True, "sources": [{
        "source_ref": "private-file-name.docx",
        "source_revision": _SOURCE_REVISION,
        "media_type": _DOCX_MEDIA,
    }]},
    {"schema": SCHEMA, "complete": True, "sources": [{
        "source_ref": _SOURCE_REF,
        "source_revision": _SOURCE_REVISION,
        "media_type": "text/plain",
        "display_name": "private name.docx",
    }]},
    {"schema": SCHEMA, "complete": True, "sources": [{
        "source_ref": _SOURCE_REF,
        "source_revision": _SOURCE_REVISION,
        "media_type": "text/plain",
        "reason_code": "private_directory_name",
    }]},
    {"schema": SCHEMA, "complete": True, "sources": [{
        "source_ref": _SOURCE_REF,
        "source_revision": "A" * 64,
        "media_type": "text/plain",
    }]},
])
async def test_list_sources_fails_closed_on_incomplete_or_unprojectable_data(payload):
    result = await ReportProviderAdapter(_SourceProvider(payload)).list_sources(_SUBJECT_REF)
    assert result == {"schema": SCHEMA, "complete": False, "sources": []}


@pytest.mark.asyncio
async def test_missing_source_listing_is_not_reported_as_an_empty_inventory():
    class NoCatalog:
        pass

    result = await ReportProviderAdapter(NoCatalog()).list_sources(_SUBJECT_REF)
    assert result == {"schema": SCHEMA, "complete": False, "sources": []}


def test_request_source_reference_bound_matches_private_catalog_bound():
    with pytest.raises(ReportContractError, match="invalid_source_refs"):
        ReportRequest.from_wire({
            "schema": SCHEMA,
            "request_id": "request-opaque-1",
            "subject_ref": _SUBJECT_REF,
            "source_refs": [_SOURCE_REF] * 201,
            "period_start": "2026-01-01",
            "period_end": "2026-01-31",
            "mode": "local",
            "template_id": "template-opaque-1",
            "mode_consent": {},
        })


class _FileArtifactStore:
    """Small persistent test store; it never reads caller-selected paths."""

    def __init__(self, root: Path):
        self.root = root
        self.root.mkdir(parents=True, exist_ok=True)

    def _directory(self, run_id: str) -> Path:
        digest = hashlib.sha256(run_id.encode("utf-8")).hexdigest()
        directory = self.root / digest
        directory.mkdir(mode=0o700, parents=True, exist_ok=True)
        return directory

    def put(self, run_id: str, content: bytes, media_type: str):
        artifact_ref = "art_" + secrets.token_hex(24)
        path = self._directory(run_id) / (artifact_ref + ".bin")
        path.write_bytes(content)
        return {"artifact_ref": artifact_ref, "content_sha256": hashlib.sha256(content).hexdigest()}

    def read(self, run_id: str, artifact_ref: str) -> bytes:
        return (self._directory(run_id) / (artifact_ref + ".bin")).read_bytes()

    def delete(self, run_id: str, artifact_ref: str) -> bool:
        path = self._directory(run_id) / (artifact_ref + ".bin")
        try:
            path.unlink()
            return True
        except FileNotFoundError:
            return False


class _FileReceiptAuthority:
    """Durable test receipt authority with exact run/source/output bindings."""

    _PUBLIC_FIELDS = (
        "verified", "coverage", "policy_version", "run_id", "source_ref",
        "source_revision", "artifact_ref", "receipt_ref", "input_sha256",
        "output_sha256",
    )

    def __init__(self, root: Path):
        self.root = root
        self.root.mkdir(parents=True, exist_ok=True)

    def issue(self, payload):
        receipt_ref = "rcpt_" + secrets.token_hex(24)
        record = {**payload, "receipt_ref": receipt_ref, "verified": True, "coverage": "complete"}
        (self.root / (receipt_ref + ".json")).write_text(
            json.dumps(record, sort_keys=True), encoding="utf-8"
        )
        return {key: record[key] for key in self._PUBLIC_FIELDS}

    def verify(self, receipt_ref, expected):
        record = json.loads((self.root / (receipt_ref + ".json")).read_text(encoding="utf-8"))
        projected = {**record, "verified": record.get("verified") is True and all(
            record.get(key) == value for key, value in expected.items()
        )}
        return {key: projected[key] for key in self._PUBLIC_FIELDS}

    def verify_report(self, report, receipt_refs):
        return {"verified": bool(receipt_refs)}


class _ProfileFactory:
    def create(self, profile_id, policy_version):
        return {
            "profile_id": profile_id,
            "policy_version": policy_version,
            "profile": {
                "client_id": "synthetic-client",
                "tarnname": "Testperson",
                "fake_geburtsdatum": "01.01.2000",
                "mappings": {"names": {"Sample Person": "Example Person"}},
                "created": "2026-10-11",
                "version": 1,
                "whitelist": [],
            },
        }


def test_document_reader_uses_generic_pipeline_and_hashes_source_revision(tmp_path):
    source = tmp_path / "synthetic.txt"
    source.write_text("Synthetic report content.", encoding="utf-8")
    revision = hashlib.sha256(source.read_bytes()).hexdigest()

    result = BachDocumentReaderPort().extract(source, expected_revision=revision, run_id="run-test")

    assert result == {
        "content": "Synthetic report content.",
        "media_type": "text/plain",
        "source_revision": revision,
        "complete": True,
    }


def test_document_reader_supports_legacy_doc_only_after_complete_generic_extraction(tmp_path):
    source = tmp_path / "synthetic.doc"
    source.write_bytes(b"synthetic legacy Word fixture")
    revision = hashlib.sha256(source.read_bytes()).hexdigest()

    class CompleteExtractor:
        def __init__(self):
            self.seen = None

        def _extract_text_from_file(self, path):
            self.seen = Path(path).suffix
            return "Complete legacy Word text"

    service = CompleteExtractor()
    result = BachDocumentReaderPort(service).extract(
        source, expected_revision=revision, run_id="run-test"
    )

    assert result == {
        "content": "Complete legacy Word text",
        "media_type": "text/plain",
        "source_revision": revision,
        "complete": True,
    }
    assert service.seen == ".doc"


def test_document_reader_rejects_legacy_doc_extraction_marker(tmp_path):
    source = tmp_path / "synthetic.doc"
    source.write_bytes(b"synthetic legacy Word fixture")
    revision = hashlib.sha256(source.read_bytes()).hexdigest()

    class FailedExtractor:
        def _extract_text_from_file(self, _path):
            return "[.doc-Extraktion fehlgeschlagen - converter unavailable]"

    with pytest.raises(ReportHostPortError) as error:
        BachDocumentReaderPort(FailedExtractor()).extract(
            source, expected_revision=revision, run_id="run-test"
        )
    assert error.value.reason_code == "source_extraction_incomplete"


def test_document_reader_supports_complete_text_only_pdf(tmp_path):
    source = tmp_path / "synthetic.pdf"
    source.write_bytes(_MINIMAL_TEXT_PDF)
    revision = hashlib.sha256(source.read_bytes()).hexdigest()

    result = BachDocumentReaderPort().extract(source, expected_revision=revision, run_id="run-test")

    assert result == {
        "content": "Probe Dokument Text",
        "media_type": "text/plain",
        "source_revision": revision,
        "complete": True,
    }


def test_document_reader_rejects_partial_pdf_with_blank_page(tmp_path):
    from io import BytesIO

    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(BytesIO(_MINIMAL_TEXT_PDF))
    writer = PdfWriter()
    writer.add_page(reader.pages[0])
    writer.add_blank_page(width=200, height=100)
    source = tmp_path / "partial.pdf"
    with source.open("wb") as target:
        writer.write(target)
    revision = hashlib.sha256(source.read_bytes()).hexdigest()

    with pytest.raises(ReportHostPortError) as error:
        BachDocumentReaderPort().extract(source, expected_revision=revision, run_id="run-test")
    assert error.value.reason_code == "source_extraction_incomplete"


def test_document_reader_rejects_unknown_format_and_changed_revision(tmp_path):
    unsupported = tmp_path / "synthetic.xls"
    unsupported.write_bytes(b"unsupported spreadsheet fixture")
    invalid_pdf = tmp_path / "invalid.pdf"
    invalid_pdf.write_bytes(b"%PDF invalid fixture")
    txt = tmp_path / "synthetic.txt"
    txt.write_text("Before", encoding="utf-8")

    with pytest.raises(ReportHostPortError) as unsupported_error:
        BachDocumentReaderPort().extract(
            unsupported, expected_revision=hashlib.sha256(unsupported.read_bytes()).hexdigest(), run_id="run-test"
        )
    with pytest.raises(ReportHostPortError) as invalid_pdf_error:
        BachDocumentReaderPort().extract(
            invalid_pdf, expected_revision=hashlib.sha256(invalid_pdf.read_bytes()).hexdigest(), run_id="run-test"
        )
    with pytest.raises(ReportHostPortError) as changed_error:
        BachDocumentReaderPort().extract(txt, expected_revision="0" * 64, run_id="run-test")

    assert getattr(unsupported_error.value, "reason_code", None) == "unsupported_source"
    assert getattr(invalid_pdf_error.value, "reason_code", None) == "source_extraction_incomplete"
    assert getattr(changed_error.value, "reason_code", None) == "source_revision_changed"


@pytest.mark.asyncio
async def test_local_model_route_without_runtime_and_slot_is_unavailable():
    from hub._services.llm.model_backend import OllamaBackend

    backend = OllamaBackend(default_model="report-local-fixture")
    local = BachModelRoutePort(backend, "report-local-fixture", "local")
    status = await local.route_status()
    assert status["available"] is False
    assert status["reason_code"] == "provider_route_unverified"
    assert status["mode"] == "local"
    assert status["local_only"] is True
    assert "route_ref" not in status
    pin = await local.pin_route(
        "run-route-unavailable", expected_mode="local", request_mode="local",
        allow_fallback=False, local_only=True,
    )
    assert pin["verified"] is False


def _local_report_runtime(
    tmp_path, monkeypatch, *, reserve_states=None, metadata_states=None, backend_chat=None,
):
    import httpx

    from hub._services import skill_source_service
    from hub._services.chat import local_resource_reserve as reserve
    from hub._services.chat import model_sockets, slots_config
    from hub._services.chat.chat_runtime import ChatRuntime
    from hub._services.llm.model_backend import OllamaBackend

    model = "report-local-fixture:small"
    slots_path = tmp_path / "slots.json"
    slots_config.initialize_slots_config(str(slots_path))
    monkeypatch.setattr(slots_config, "DEFAULT_SLOTS_FILE", str(slots_path))
    monkeypatch.setenv("BACH_SLOTS_CONFIG_PATH", str(slots_path))
    monkeypatch.setenv("BACH_RUNTIME_DIR", str(tmp_path / "runtime"))
    monkeypatch.setattr(model_sockets.platform, "node", lambda: "report-test-host")
    monkeypatch.setattr(skill_source_service, "check_write_locks", lambda _target: None)

    config = slots_config.load_slots_config(str(slots_path), strict=True)
    profile = config["slots"]["buddha_chat"]
    profile.update({
        "enabled": True,
        "backend": "ollama",
        "model": model,
        "skill_refs": [],
        "model_fallbacks": [],
    })
    slots_config.save_slots_config(config, str(slots_path))
    version = slots_config.core_system_agents_snapshot(str(slots_path))["configuration_version"]
    model_sockets.migrate_model_sockets(version, path=str(slots_path))

    backend = OllamaBackend(base_url="http://127.0.0.1:11434", default_model=model)
    runtime = ChatRuntime(backend)
    runtime.require_model_socket_config = True
    delegated = []

    async def default_chat(messages, *, tools, model):
        from hub._services.chat.host_inference_gate import HostInferenceGate

        delegated.append({
            "messages": messages,
            "tools": tools,
            "model": model,
            "gate": HostInferenceGate().status(),
        })
        return {"content": "synthetic local response"}

    monkeypatch.setattr(backend, "chat", backend_chat or default_chat)
    metadata_states = iter(metadata_states or [])

    class ShowResponse:
        def __init__(self, payload):
            self.payload = payload

        def raise_for_status(self):
            return None

        def json(self):
            return self.payload

    def ollama_show(url, **kwargs):
        assert url == "http://127.0.0.1:11434/api/show"
        assert kwargs.get("follow_redirects") is False
        state = next(metadata_states, "local")
        payload = {
            "details": {"format": "gguf"},
            "model_info": {"general.parameter_count": 123},
        }
        if state == "remote":
            payload.update({"remote_host": "https://cloud.example.invalid", "remote_model": model})
        return ShowResponse(payload)

    monkeypatch.setattr(httpx, "post", ollama_show)
    policy_version = "c" * 64
    monkeypatch.setattr(reserve, "read_policy", lambda: ({"enabled": True}, policy_version))
    states = list(reserve_states or ["ready", "ready", "ready"])
    probe_count = []

    async def probe(_backend, probe_model):
        probe_count.append(probe_model)
        state = states.pop(0)
        if isinstance(state, dict):
            return state
        return {
            "schema": reserve.SCHEMA,
            "state": state,
            "admitted": state == "ready",
            "enforced": state != "not_configured",
            "retryable": False,
            "policy_version": policy_version,
            "model": probe_model,
            "wait_seconds": 1,
            "poll_seconds": 1,
        }

    monkeypatch.setattr(reserve, "probe", probe)
    local = BachModelRoutePort(
        backend, model, "local", runtime=runtime, agent_id="buddha_chat",
    )
    return local, runtime, backend, delegated, probe_count, states


@pytest.mark.asyncio
async def test_local_report_dispatch_uses_runtime_gates_and_foreground_slot(tmp_path, monkeypatch):
    from hub._services.chat.host_inference_gate import HostInferenceGate

    local, runtime, _backend, delegated, probes, states = _local_report_runtime(tmp_path, monkeypatch)
    status = await local.route_status()
    assert status["available"] is True
    assert status["mode"] == "local"
    assert status["local_only"] is True
    assert status["route_id"].startswith("local_")
    assert "route_ref" not in status
    pin = await local.pin_route(
        "run-route-runtime", expected_mode="local", request_mode="local",
        allow_fallback=False, local_only=True,
    )
    assert pin["verified"] is True
    response = await local.generate(
        "synthetic prompt", run_id="run-route-runtime", mode="local", request_mode="local",
        route_id=pin["route_id"], route_ref=pin["route_ref"],
        allow_fallback=False, local_only=True,
    )
    assert response["text"] == "synthetic local response"
    assert len(delegated) == 1 and delegated[0]["tools"] is None
    assert delegated[0]["gate"]["active"] is True
    assert delegated[0]["gate"]["priority"] == "foreground"
    assert delegated[0]["gate"]["model_target"]["agent_id"] == "buddha_chat"
    assert len(probes) == 3 and not states
    assert not runtime.sessions
    assert runtime._compute_turn_context.get() is None
    assert not runtime._session_locks and not runtime._chat_turn_gates
    assert HostInferenceGate().status()["active"] is False


@pytest.mark.asyncio
async def test_local_report_reserve_denial_cleans_session_without_backend_call(tmp_path, monkeypatch):
    local, runtime, _backend, delegated, probes, states = _local_report_runtime(
        tmp_path, monkeypatch, reserve_states=["not_configured"],
    )
    pin = await local.pin_route(
        "run-route-denied", expected_mode="local", request_mode="local",
        allow_fallback=False, local_only=True,
    )
    assert pin["verified"] is True
    with pytest.raises(ReportHostPortError):
        await local.generate(
            "synthetic prompt", run_id="run-route-denied", mode="local", request_mode="local",
            route_id=pin["route_id"], route_ref=pin["route_ref"],
            allow_fallback=False, local_only=True,
        )
    assert not delegated and len(probes) == 1 and not states
    assert not runtime.sessions and runtime._compute_turn_context.get() is None
    assert not runtime._session_locks and not runtime._chat_turn_gates


@pytest.mark.asyncio
async def test_external_target_dispatch_after_gate_release_is_rejected(tmp_path, monkeypatch):
    local, runtime, _backend, delegated, probes, states = _local_report_runtime(
        tmp_path, monkeypatch, reserve_states=["ready", "external_target"],
    )
    pin = await local.pin_route(
        "run-route-alias", expected_mode="local", request_mode="local",
        allow_fallback=False, local_only=True,
    )
    assert pin["verified"] is True
    with pytest.raises(ReportHostPortError):
        await local.generate(
            "synthetic prompt", run_id="run-route-alias", mode="local", request_mode="local",
            route_id=pin["route_id"], route_ref=pin["route_ref"],
            allow_fallback=False, local_only=True,
        )
    assert not delegated and len(probes) == 2 and not states
    assert not runtime.sessions and runtime._compute_turn_context.get() is None
    assert not runtime._session_locks and not runtime._chat_turn_gates


@pytest.mark.asyncio
async def test_local_metadata_remote_change_at_dispatch_is_rejected(tmp_path, monkeypatch):
    local, runtime, _backend, delegated, probes, states = _local_report_runtime(
        tmp_path, monkeypatch,
        metadata_states=["local"] * 4 + ["remote"],
    )
    pin = await local.pin_route(
        "run-route-remote-metadata", expected_mode="local", request_mode="local",
        allow_fallback=False, local_only=True,
    )
    assert pin["verified"] is True
    with pytest.raises(ReportHostPortError) as error:
        await local.generate(
            "synthetic prompt", run_id="run-route-remote-metadata", mode="local", request_mode="local",
            route_id=pin["route_id"], route_ref=pin["route_ref"],
            allow_fallback=False, local_only=True,
        )
    assert error.value.reason_code == "provider_route_unverified"
    assert not delegated and len(probes) == 3 and not states
    assert not runtime.sessions and runtime._compute_turn_context.get() is None
    assert not runtime._session_locks and not runtime._chat_turn_gates


def test_native_local_model_metadata_rejects_remote_ollama_alias(monkeypatch):
    import httpx

    from hub._services.agent_manage_service import verified_local_agent_model
    from hub._services.llm.model_backend import OllamaBackend

    backend = OllamaBackend(base_url="http://127.0.0.1:11434", default_model="alias:small")

    class Response:
        def __init__(self, payload):
            self.payload = payload

        def raise_for_status(self):
            return None

        def json(self):
            return self.payload

    payload = {
        "details": {"format": "gguf"},
        "model_info": {"general.parameter_count": 123},
        "remote_host": "https://cloud.example.invalid",
        "remote_model": "alias:small",
    }
    calls = []

    def post(url, **kwargs):
        calls.append((url, kwargs))
        return Response(payload)

    monkeypatch.setattr(httpx, "post", post)
    assert verified_local_agent_model(backend, "alias:small") is False
    assert len(calls) == 1
    assert calls[0][0] == "http://127.0.0.1:11434/api/show"
    assert calls[0][1]["follow_redirects"] is False


@pytest.mark.asyncio
async def test_local_host_port_without_trusted_runtime_binding_is_unavailable():
    from hub._services.llm.model_backend import OllamaBackend

    backend = OllamaBackend(default_model="report-local-fixture:small")
    ports = build_report_host_ports(local_model={
        "backend": backend,
        "model": "report-local-fixture:small",
    })
    status = await ports["local_model"].route_status()
    assert status["available"] is False
    assert status["reason_code"] == "provider_route_unverified"
    assert status["verified_local"] is False


@pytest.mark.asyncio
async def test_local_backend_exception_cleans_ephemeral_session(tmp_path, monkeypatch):
    async def failing_chat(_messages, *, tools, model):
        assert tools is None
        raise RuntimeError("synthetic backend exception")

    local, runtime, _backend, delegated, _probes, _states = _local_report_runtime(
        tmp_path, monkeypatch, backend_chat=failing_chat,
    )
    pin = await local.pin_route(
        "run-route-exception", expected_mode="local", request_mode="local",
        allow_fallback=False, local_only=True,
    )
    with pytest.raises(ReportHostPortError):
        await local.generate(
            "synthetic prompt", run_id="run-route-exception", mode="local", request_mode="local",
            route_id=pin["route_id"], route_ref=pin["route_ref"],
            allow_fallback=False, local_only=True,
        )
    assert not runtime.sessions and runtime._compute_turn_context.get() is None
    assert not runtime._session_locks and not runtime._chat_turn_gates


@pytest.mark.asyncio
async def test_cloud_route_status_has_explicit_nonlocal_mode_and_no_fallback():
    class Backend:
        async def chat(self, _messages, *, tools, model):
            assert tools is None
            return {"content": "synthetic response"}

    async def attest(_backend, _model):
        return {
            "verified": True,
            "route_id": "cloud_attested_route_1",
            "cloud_route": True,
            "fallback_disabled": True,
        }

    cloud = BachModelRoutePort(Backend(), "cloud-model", "cloud", cloud_attestor=attest)
    status = await cloud.route_status()
    assert status["available"] is True
    assert status["mode"] == "cloud"
    assert status["local_only"] is False
    assert status["fallback_disabled"] is True
    assert "route_ref" not in status


@pytest.mark.asyncio
async def test_word_template_port_uses_real_service_and_bound_resolver(tmp_path):
    from docx import Document

    from hub._services.document.word_template_service import WordTemplateService

    template_path = tmp_path / "template.docx"
    document = Document()
    document.add_paragraph("{{REPORT_TITLE}}")
    document.save(template_path)
    template_bytes = template_path.read_bytes()
    store = _FileArtifactStore(tmp_path / "artifacts")

    class Resolver:
        def resolve(self, template_id, *, expected_revision):
            assert template_id == "template-public"
            assert expected_revision == _TEMPLATE_REVISION
            return {
                "available": True,
                "revision": _TEMPLATE_REVISION,
                "media_type": _DOCX_MEDIA,
                "content": template_bytes,
            }

    port = BachWordTemplatePort(store, service=WordTemplateService())
    port.bind_resolver(Resolver())
    rendered = await port.render_draft(
        "foerderplaner.report.v1",
        {"REPORT_TITLE": "Synthetic report"},
        "template-public",
        _TEMPLATE_REVISION,
        run_id="run-test",
    )
    content = store.read("run-test", rendered["artifact_ref"])
    output_path = tmp_path / "rendered.docx"
    output_path.write_bytes(content)
    result_doc = Document(output_path)

    assert rendered["template_revision"] == _TEMPLATE_REVISION
    assert "Synthetic report" in "\n".join(paragraph.text for paragraph in result_doc.paragraphs)


@pytest.mark.asyncio
async def test_real_document_anonymizer_requires_and_rechecks_bound_receipt(tmp_path):
    source_content = "Sample Person completed a synthetic assessment.".encode("utf-8")
    source_sha = hashlib.sha256(source_content).hexdigest()
    store = _FileArtifactStore(tmp_path / "artifacts")
    receipts = _FileReceiptAuthority(tmp_path / "receipts")
    port = BachDocumentAnonymizerPort(
        store,
        _ProfileFactory(),
        receipts,
    )

    anonymized = await port.anonymize({
        "source_ref": _SOURCE_REF,
        "source_revision": source_sha,
        "media_type": "text/plain",
        "content": source_content,
        "anonymization_profile": {
            "profile_id": "profile_" + "3" * 32,
            "policy_version": "policy_1",
        },
    }, run_id="run-test")
    assert anonymized["status"] == "ok"
    assert "Sample Person" not in store.read("run-test", anonymized["artifact_ref"]).decode("utf-8")
    assert "Example Person" in store.read("run-test", anonymized["artifact_ref"]).decode("utf-8")

    early_read = None
    with pytest.raises(ReportHostPortError) as before_receipt_error:
        await port.read_artifact(anonymized["artifact_ref"], run_id="run-test")
    early_read = getattr(before_receipt_error.value, "reason_code", None)

    input_sha = hashlib.sha256(source_content).hexdigest()
    verified = await port.verify_receipt(
        anonymized["receipt_ref"],
        run_id="run-test",
        source_ref=_SOURCE_REF,
        source_revision=source_sha,
        artifact_ref=anonymized["artifact_ref"],
        expected_input_sha256=input_sha,
    )
    readback = await port.read_artifact(anonymized["artifact_ref"], run_id="run-test")
    assert verified["verified"] is True
    assert verified["coverage"] == "complete"
    assert verified["source_ref"] == _SOURCE_REF
    assert verified["source_revision"] == source_sha
    assert verified["input_sha256"] == input_sha
    assert verified["output_sha256"] == hashlib.sha256(readback).hexdigest()
    assert early_read == "receipt_required"


@pytest.mark.asyncio
async def test_anonymizer_run_cleanup_reports_partial_failure(tmp_path):
    class PartialStore:
        async def delete(self, run_id, artifact_ref):
            return artifact_ref == "art_" + "4" * 48

    port = BachDocumentAnonymizerPort(
        PartialStore(), _ProfileFactory(), _FileReceiptAuthority(tmp_path / "receipts")
    )
    result = await port.delete_run_artifacts(
        "run-test", ["art_" + "4" * 48, "art_" + "5" * 48]
    )
    assert result == {"complete": False, "deleted_count": 1}


def test_host_port_builder_keeps_document_reader_available_and_unconfigured_models_absent():
    ports = build_report_host_ports()
    assert isinstance(ports["document_reader"], BachDocumentReaderPort)
    assert "local_model" not in ports
    assert "cloud_model" not in ports


@pytest.mark.asyncio
async def test_entry_point_boot_uses_shared_host_ports_without_installing_provider(monkeypatch):
    from hub._services.document import report_provider_adapter as module

    class Provider:
        def __init__(self, host_ports):
            self.host_ports = host_ports

        async def capabilities(self):
            return _CAPABILITIES

        async def prepare(self, request):
            raise AssertionError("boot test does not prepare a report")

        async def finish(self, run_id, approval):
            raise AssertionError("boot test does not finish a report")

        async def status(self, run_id):
            raise AssertionError("boot test does not query run status")

        async def cleanup(self, run_id, approval):
            raise AssertionError("boot test does not clean up")

        async def list_sources(self, subject_ref):
            return {"schema": SCHEMA, "complete": True, "sources": []}

    provider = Provider

    class Entry:
        name = "foerderplaner"

        def load(self):
            return lambda host_ports: provider(host_ports)

    class Entries(list):
        def select(self, *, group):
            assert group == "ellmos.report_providers"
            return self

    monkeypatch.setattr(module.importlib.metadata, "entry_points", lambda: Entries([Entry()]))
    host_ports = build_report_host_ports()
    adapter = load_report_provider(host_ports, provider_name="foerderplaner")

    assert (await adapter.capabilities())["available"] is True
    assert (await adapter.capabilities())["mode_statuses"]["cloud"]["reason_code"] == "cloud_consent_unavailable"
    assert isinstance(adapter._provider.host_ports["document_reader"], BachDocumentReaderPort)
    assert await adapter.list_sources(_SUBJECT_REF) == {
        "schema": SCHEMA,
        "complete": True,
        "sources": [],
    }


@pytest.mark.asyncio
async def test_capabilities_invalid_projection_fails_closed_without_private_details():
    class Provider:
        async def capabilities(self):
            return {
                **_CAPABILITIES,
                "mode_statuses": {
                    "local": {"available": True},
                    "cloud": {"available": False, "reason_code": "private_directory_name"},
                    "hybrid": {"available": False, "reason_code": "hybrid_policy_unavailable"},
                },
            }

    result = await ReportProviderAdapter(Provider()).capabilities()

    assert result["available"] is False
    assert result["modes"] == []
    assert result["reason_code"] == "provider_unavailable"
    assert all(not status["available"] for status in result["mode_statuses"].values())
    assert "private_directory_name" not in json.dumps(result)


@pytest.mark.asyncio
async def test_non_neutral_reason_codes_are_redacted_for_finish_status_cleanup_and_watcher():
    class Provider:
        async def finish(self, run_id, approval):
            return {
                "schema": SCHEMA,
                "run_id": run_id,
                "status": "blocked",
                "effect": "none",
                "reason_code": "private_directory_name",
                "artifact_refs": [],
                "review_required": True,
            }

        async def status(self, run_id):
            return {
                "schema": SCHEMA,
                "run_id": run_id,
                "status": "prepared",
                "review_required": True,
                "reason_code": "private_directory_name",
            }

        async def cleanup(self, run_id, approval):
            return {
                "schema": SCHEMA,
                "run_id": run_id,
                "status": "cleaned",
                "sources_retained": True,
                "reason_code": "private_directory_name",
            }

        async def watcher_status(self):
            return {
                "schema": SCHEMA,
                "status": "disabled",
                "enabled": False,
                "reason_code": "private_directory_name",
            }

    adapter = ReportProviderAdapter(Provider())
    run_id = "run-opaque-1"
    with pytest.raises(ReportContractError, match="invalid_reason_code"):
        await adapter.finish(run_id, {
            "approved": True,
            "preview_ref": "preview-opaque-1",
            "consent_ref": "consent-opaque-1",
        })
    status = await adapter.status(run_id)
    cleanup = await adapter.cleanup(run_id, {
        "run_id": run_id,
        "consent_ref": "consent-opaque-1",
        "approved": True,
        "delete_sources": False,
    })
    watcher = await adapter.watcher_status()

    assert status["reason_code"] == "provider_unavailable"
    assert cleanup["reason_code"] == "cleanup_failed"
    assert cleanup["sources_retained"] is True
    assert watcher["reason_code"] == "watcher_unavailable"
    assert "private_directory_name" not in json.dumps({"status": status, "cleanup": cleanup, "watcher": watcher})


@pytest.mark.asyncio
async def test_exact_private_facade_reason_codes_survive_each_response_projection():
    class Provider:
        async def capabilities(self):
            return {
                **_CAPABILITIES,
                "mode_statuses": {
                    "local": {"available": False, "reason_code": "provider_route_unverified"},
                    "cloud": {"available": False, "reason_code": "provider_consent_denied"},
                    "hybrid": {"available": False, "reason_code": "provider_hybrid_schema_invalid"},
                },
                "modes": [],
                "available": False,
                "reason_code": "provider_not_configured",
            }

        async def prepare(self, request):
            return {
                "schema": SCHEMA,
                "run_id": None,
                "status": "blocked",
                "preview_ref": None,
                "expires_at": None,
                "request_digest": request["mode_consent"]["request_digest"],
                "source_revision": None,
                "template_revision": None,
                "source_count": 0,
                "reason_code": "provider_source_unavailable",
            }

        async def finish(self, run_id, approval):
            return {
                "schema": SCHEMA,
                "run_id": run_id,
                "status": "blocked",
                "effect": "none",
                "reason_code": "provider_finish_denied",
                "artifact_refs": [],
                "review_required": True,
            }

        async def status(self, run_id):
            return {
                "schema": SCHEMA,
                "run_id": run_id,
                "status": "blocked",
                "review_required": True,
                "reason_code": "provider_run_not_found",
            }

        async def cleanup(self, run_id, approval):
            return {
                "schema": SCHEMA,
                "run_id": run_id,
                "status": "cleanup_failed",
                "sources_retained": True,
                "reason_code": "provider_cleanup_partial",
            }

        async def watcher_status(self):
            return {
                "schema": SCHEMA,
                "status": "disabled",
                "enabled": False,
                "reason_code": "provider_watcher_disabled",
            }

    adapter = ReportProviderAdapter(Provider())
    capabilities = await adapter.capabilities()
    assert capabilities["reason_code"] == "provider_not_configured"
    assert capabilities["mode_statuses"]["local"]["reason_code"] == "provider_route_unverified"

    prepared = await adapter.prepare(_request_wire())
    assert prepared["reason_code"] == "provider_source_unavailable"
    finished = await adapter.finish("run-opaque-1", {
        "approved": True,
        "preview_ref": "preview-opaque-1",
        "consent_ref": "consent-opaque-1",
    })
    assert finished["reason_code"] == "provider_finish_denied"
    assert (await adapter.status("run-opaque-1"))["reason_code"] == "provider_run_not_found"
    assert (await adapter.cleanup("run-opaque-1", {
        "run_id": "run-opaque-1", "consent_ref": "consent-opaque-1", "approved": True,
    }))["reason_code"] == "provider_cleanup_partial"
    assert (await adapter.watcher_status())["reason_code"] == "provider_watcher_disabled"


@pytest.mark.asyncio
async def test_watcher_config_projects_explicit_subject_scope_only():
    class Provider:
        def __init__(self):
            self.config = None

        async def watcher_configure(self, config):
            self.config = config
            return {"schema": SCHEMA, "status": "disabled", "enabled": False}

    provider = Provider()
    adapter = ReportProviderAdapter(provider)
    config = _watcher_wire(subject_refs=[_SUBJECT_REF])

    response = await adapter.watcher_configure(config)

    assert provider.config["subject_refs"] == [_SUBJECT_REF]
    assert provider.config == WatcherConfig.from_wire(config).to_wire()
    assert provider.config["schema"] == SCHEMA
    assert response == {"schema": SCHEMA, "status": "disabled", "enabled": False}


@pytest.mark.asyncio
async def test_watcher_rejects_empty_enabled_scope_before_private_call():
    class Provider:
        def __init__(self):
            self.called = False

        async def watcher_configure(self, config):
            self.called = True
            return {"schema": SCHEMA, "status": "enabled", "enabled": True}

    provider = Provider()
    with pytest.raises(ReportContractError, match="invalid_watcher_subject_refs"):
        await ReportProviderAdapter(provider).watcher_configure(
            _watcher_wire(enabled=True, subject_refs=[])
        )
    assert provider.called is False


@pytest.mark.asyncio
async def test_cleanup_source_deletion_is_blocked_before_private_effects():
    class Provider:
        def __init__(self):
            self.called = False

        async def cleanup(self, run_id, approval):
            self.called = True
            raise AssertionError("source deletion must be blocked at the public boundary")

    provider = Provider()
    result = await ReportProviderAdapter(provider).cleanup(
        "run-opaque-1",
        {
            "run_id": "run-opaque-1",
            "consent_ref": "consent-opaque-1",
            "approved": True,
            "delete_sources": True,
        },
    )

    assert result == {
        "schema": SCHEMA,
        "run_id": "run-opaque-1",
        "status": "blocked",
        "reason_code": "source_deletion_unsupported",
        "sources_retained": True,
    }
    assert provider.called is False


@pytest.mark.asyncio
async def test_cleanup_preserves_partial_failure_and_retention_status():
    class Provider:
        async def cleanup(self, run_id, approval):
            assert approval["delete_sources"] is False
            return {
                "schema": SCHEMA,
                "run_id": run_id,
                "status": "cleanup_failed",
                "reason_code": "artifact_cleanup_partial",
                "sources_retained": False,
            }

    result = await ReportProviderAdapter(Provider()).cleanup(
        "run-opaque-2",
        {
            "run_id": "run-opaque-2",
            "consent_ref": "consent-opaque-2",
            "approved": True,
            "delete_sources": False,
        },
    )

    assert result["status"] == "cleanup_failed"
    assert result["sources_retained"] is False
    assert result["reason_code"] == "artifact_cleanup_partial"


@pytest.mark.asyncio
async def test_status_cleanup_and_watcher_reject_extra_private_fields():
    class Provider:
        async def status(self, run_id):
            return {
                "schema": SCHEMA,
                "run_id": run_id,
                "status": "prepared",
                "review_required": True,
                "private_path": "not-for-wire",
            }

        async def cleanup(self, run_id, approval):
            return {
                "schema": SCHEMA,
                "run_id": run_id,
                "status": "cleaned",
                "sources_retained": True,
                "private_template": "not-for-wire",
            }

        async def watcher_status(self):
            return {
                "schema": SCHEMA,
                "status": "disabled",
                "enabled": False,
                "private_root": "not-for-wire",
            }

    adapter = ReportProviderAdapter(Provider())
    run_id = "run-opaque-3"
    status = await adapter.status(run_id)
    cleanup = await adapter.cleanup(run_id, {
        "run_id": run_id,
        "consent_ref": "consent-opaque-3",
        "approved": True,
    })
    watcher = await adapter.watcher_status()

    assert status["status"] == "unknown"
    assert cleanup["status"] == "unknown"
    assert watcher["status"] == "unknown"
    assert "not-for-wire" not in json.dumps({"status": status, "cleanup": cleanup, "watcher": watcher})


@pytest.mark.asyncio
async def test_operator_confirmation_is_verification_only_and_binding_exact():
    import hashlib
    from datetime import datetime, timedelta, timezone

    request_expiry = (datetime.now(timezone.utc) + timedelta(minutes=10)).isoformat()
    proof_expiry = (datetime.now(timezone.utc) + timedelta(minutes=20)).isoformat()
    binding = {
        "action": "report.mode_consent",
        "request_id": "request-opaque-1",
        "request_digest": hashlib.sha256(b"stable request").hexdigest(),
        "subject_ref": _SUBJECT_REF,
        "mode": "local",
        "template_id": "template-opaque-1",
        "expires_at": request_expiry,
    }

    class Verifier:
        def __init__(self):
            self.calls = []

        async def verify_confirmation(self, confirmation_ref, expected_binding):
            self.calls.append((confirmation_ref, expected_binding))
            return {
                "verified": True,
                "operator_confirmed": True,
                "binding": expected_binding,
                "proof_ref": "proof-opaque-1",
                "expires_at": proof_expiry,
            }

    verifier = Verifier()
    port = BachOperatorConfirmationPort(verifier)
    proof = await port.verify_confirmation("confirmation-opaque-1", binding)

    assert proof["verified"] is True
    assert proof["operator_confirmed"] is True
    assert proof["binding"]["expires_at"].endswith("Z")
    assert proof["expires_at"].endswith("Z")
    assert verifier.calls[0][0] == "confirmation-opaque-1"
    assert not hasattr(port, "issue_consent")


@pytest.mark.asyncio
@pytest.mark.parametrize("proof_update", [
    {"operator_confirmed": False},
    {"binding": {"action": "report.finish"}},
    {"expires_at": "2099-01-01T00:00:00"},
])
async def test_operator_confirmation_rejects_untrusted_or_malformed_proofs(proof_update):
    from datetime import datetime, timedelta, timezone

    expiry = (datetime.now(timezone.utc) + timedelta(minutes=10)).isoformat()
    binding = {
        "action": "report.watcher.configure",
        "config_digest": "a" * 64,
        "expires_at": expiry,
    }

    class Verifier:
        async def verify_confirmation(self, confirmation_ref, expected_binding):
            return {
                "verified": True,
                "operator_confirmed": True,
                "binding": expected_binding,
                "proof_ref": "proof-opaque-2",
                "expires_at": expiry,
                **proof_update,
            }

    with pytest.raises(ReportHostPortError):
        await BachOperatorConfirmationPort(Verifier()).verify_confirmation("confirmation-opaque-2", binding)
